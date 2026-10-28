import json
import os

import pytest

from import_worker import dossier_builder as b
from import_worker import normalize as n
from import_worker.parse_dossier import DossierFehler, normalize_dossier, parse_docx
from import_worker.slug import check_slug

from .conftest import SEED_DIR, seed_paths


def _roundtrip(seed_path, tmp_path, opts=None):
    folder = b.build_package(seed_path, str(tmp_path), opts)
    doc = [f for f in os.listdir(folder) if f.endswith(".docx")][0]
    d = parse_docx(os.path.join(folder, doc))
    return d, normalize_dossier(d)


def test_pressel_felder(pressel_seed, tmp_path):
    d, dr = _roundtrip(os.path.join(SEED_DIR, "naturlehrpfad-pressel.json"), tmp_path)
    t = dr["trail"]
    assert d["template_version"] == 3
    assert t["id"] == "naturlehrpfad-pressel"
    assert t["typ"] == "wald" and t["form"] == "linie" and t["land"] == "DE"
    assert t["laenge_km"] == pressel_seed["laengeKm"]
    assert t["dauer_min"] == pressel_seed["dauerMin"]
    assert t["rundkurs"] is True
    assert t["tags"] == pressel_seed["tags"]
    assert t["lebensraeume"] == pressel_seed["lebensraeume"]
    assert t["naturraum"] == pressel_seed["naturraum"]
    assert t["region"] == pressel_seed["region"]
    assert abs(t["start_lat"] - pressel_seed["route"][0][0]) < 1e-5
    assert dr["hinweise"] == []
    assert dr["meta"]["geometrie_datei"] == "route.gpx"
    assert dr["meta"]["geometrie_lizenz"] == "eigen"


def test_stationen_und_arten(pressel_seed, tmp_path):
    _, dr = _roundtrip(os.path.join(SEED_DIR, "naturlehrpfad-pressel.json"), tmp_path)
    assert len(dr["stationen"]) == len(pressel_seed["stationen"])
    s0, e0 = dr["stationen"][0], pressel_seed["stationen"][0]
    for k in ("reihenfolge", "km", "lat", "lon", "titel", "thema", "kurztext", "erlebnisse"):
        assert s0[k] == e0[k]
    assert [a["name_de"] for a in dr["arten"]] == pressel_seed["arten"]
    assert dr["arten"][0]["nachweis"] == "belegt"
    assert dr["arten"][0]["name_lat"] == "Grus grus"


def test_quellen_je_abschnitt(tmp_path):
    d, dr = _roundtrip(os.path.join(SEED_DIR, "naturlehrpfad-pressel.json"), tmp_path)
    assert dr["meta"]["quellen_je_abschnitt"]["1"] == 1
    assert dr["quellen"][0]["abgerufen_am"] == "2026-10-22"


def test_leere_vorlage_ist_leer(tmp_path):
    d = parse_docx(b.TEMPLATE)
    dr = normalize_dossier(d)
    assert dr["stationen"] == [] and dr["arten"] == [] and dr["bilder"] == []
    assert dr["trail"]["name"] is None
    assert dr["meta"]["abschnitte"]["1"] == "offen"
    assert dr["meta"]["gefuellte_felder"] == []


def test_kein_dossier(tmp_path):
    p = tmp_path / "x.docx"
    p.write_bytes(b"kein zip")
    with pytest.raises(DossierFehler):
        parse_docx(str(p))


def test_v2_vorlage_land_nicht_region(tmp_path):
    """v2: 'neu, Spalte geplant (bis dahin trails.region)' darf Region nicht ueberschreiben."""
    import docx
    v2 = os.path.expanduser("~/Downloads/Vorlage_Pfad-Dossier_v2.docx")
    if not os.path.isfile(v2):
        pytest.skip("v2-Vorlage nicht vorhanden")
    doc = docx.Document(v2)
    for r in doc.tables[0].rows:
        if r.cells[0].text.startswith("Land"):
            r.cells[2].paragraphs[0].add_run("DE")
        if r.cells[0].text.startswith("Region"):
            r.cells[2].paragraphs[0].add_run("Prignitz")
    p = tmp_path / "v2.docx"
    doc.save(str(p))
    dr = normalize_dossier(parse_docx(str(p)))
    assert dr["trail"]["land"] == "DE" and dr["trail"]["region"] == "Prignitz"


@pytest.mark.parametrize("path", seed_paths(), ids=lambda p: os.path.basename(p)[:-5])
def test_alle_seeds_roundtrip(path, tmp_path):
    with open(path, encoding="utf-8") as f:
        seed = json.load(f)
    if not (seed.get("route") or seed.get("area")):
        pytest.skip("kein Trail-Seed")
    _, dr = _roundtrip(path, tmp_path)
    t = dr["trail"]
    assert t["id"] == seed["id"]
    assert t["name"] == seed["name"]
    assert t["typ"] == seed["typ"]
    assert t["kurzbeschreibung"] == seed["kurzbeschreibung"].strip()
    assert len(dr["stationen"]) == len(seed["stationen"])
    assert [s["reihenfolge"] for s in dr["stationen"]] == [s["reihenfolge"] for s in seed["stationen"]]
    assert t["lebensraeume"] == (seed.get("lebensraeume") or [])
    assert not [h for h in dr["hinweise"] if h["code"] == "SCHEMA"], dr["hinweise"]


def test_normalize_helfer():
    assert n.to_float("5,4 km") == (5.4, True)
    assert n.to_float("5.4") == (5.4, False)
    assert n.to_float("abc") == (None, False)
    assert n.to_int("ca. 90") == (90, True)
    assert n.to_bool("Ja") == (True, True)
    assert n.to_bool("nein") == (False, False)
    assert n.to_bool("vielleicht") == (None, False)
    assert n.split_latlon("53.5512, 13.2471")[:2] == (53.5512, 13.2471)
    assert n.split_latlon("53,5512 / 13,2471")[:2] == (53.5512, 13.2471)
    assert n.to_list("a, b;c") == ["a", "b", "c"]
    assert n.to_date("22.10.2026") == "2026-10-22"


def test_slug():
    assert check_slug("kollhorst") == ("kollhorst", False, True)
    assert check_slug("Märkisch Buchholz") == ("maerkisch-buchholz", True, True)
    assert check_slug("") == ("", False, False)
    assert check_slug("!!!")[2] is False
