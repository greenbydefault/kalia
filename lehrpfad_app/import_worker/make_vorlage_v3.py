#!/usr/bin/env python3
"""Erzeugt Vorlage_Pfad-Dossier_v3.docx aus der v2-Vorlage.

Aenderungen gegenueber v2 (siehe VORLAGE_V3.md):
  - template_version: 3
  - Land/Bundesland/Kreis/Ort/Schutzgebiet zeigen auf trails.land ... (Spalten ab 0035)
  - Abschnitt 3: Geometrie-Quelle und -Lizenz (Pflicht, wenn route.gpx/flaeche.geojson beiliegt)
  - Abschnitt 6: Steckbrief-Tabelle statt Freitext
  - Abschnitt 7: Geo-Check-Spalte entfaellt (der Import liest EXIF selbst)
  - Anhang A: A3-A6 vollstaendig, neu A8 (Geometrie-Lizenzen)

Nutzung:
  python3 -m import_worker.make_vorlage_v3 <v2.docx> [out.docx]
"""
from __future__ import annotations

import copy
import os
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from . import lists
from .docxutil import add_row_like_last, drop_col, set_cell, set_text

DEFAULT_OUT = os.path.join(os.path.dirname(__file__), "templates", "Vorlage_Pfad-Dossier_v3.docx")


def _para_after(doc, anchor_text_start: str):
    for p in doc.paragraphs:
        if p.text.startswith(anchor_text_start):
            return p
    raise SystemExit("Anker nicht gefunden: %s" % anchor_text_start)


def _following(doc, paragraph):
    el = paragraph._p.getnext()
    return Paragraph(el, doc)


def build(src: str, out: str) -> None:
    doc = docx.Document(src)

    set_text(_para_after(doc, "template_version:"), "template_version: 3 · Stand 22.10.2026")

    # --- Abschnitt 1: neue Spalten
    sec1 = doc.tables[0]
    spalten = {
        "Land": "trails.land", "Bundesland": "trails.bundesland",
        "Kreis": "trails.kreis", "Ort": "trails.ort", "Schutzgebiet": "trails.schutzgebiet",
    }
    for row in sec1.rows[1:]:
        label = row.cells[0].text
        for key, col in spalten.items():
            if label.startswith(key):
                set_cell(row.cells[1], col)
    for row in sec1.rows[1:]:
        if row.cells[0].text.startswith("Region (Freitext"):
            set_cell(row.cells[0], "Region (Freitext; leer = Ort, Kreis, Bundesland)")

    # --- Abschnitt 3: Geometrie-Herkunft
    sec3 = doc.tables[4]
    for label, col in (
        ("Geometrie-Quelle (Name oder URL der GPX/GeoJSON-Datei; leer bei OSM-Pipeline)", "import.geometrie_quelle"),
        ("Geometrie-Lizenz (Key aus Anhang A8; Pflicht, wenn eine Datei beiliegt)", "import.geometrie_lizenz"),
    ):
        row = add_row_like_last(sec3)
        set_cell(row.cells[0], label)
        set_cell(row.cells[1], col)
        set_cell(row.cells[2], "")

    # --- Abschnitt 6: Steckbrief-Tabelle (Kopie einer Quellen-Tabelle mit 3 Spalten)
    steck_anchor = _para_after(doc, "Steckbrief (optional, pro Station)")
    set_text(steck_anchor, "Steckbrief (optional, pro Station): eine Zeile je Schlüssel. Keys: "
             + ", ".join(lists.STECKBRIEF_KEYS) + ".")
    quellen = doc.tables[1]._tbl
    steck = copy.deepcopy(quellen)
    steck_anchor._p.addnext(steck)
    st = Table(steck, doc)
    drop_col(st, 3)
    for cell, text in zip(st.rows[0].cells, ("Station (Reihenfolge)", "Schlüssel", "Wert")):
        set_cell(cell, text)
    for cell in st.rows[1].cells:
        set_cell(cell, "")

    # --- Abschnitt 7: Geo-Check raus
    bilder = next(t for t in doc.tables if t.rows[0].cells[0].text.startswith("Sort") and len(t.columns) >= 10)
    assert bilder.rows[0].cells[-1].text.startswith("Geo-Check"), "Bilder-Tabelle erwartet"
    drop_col(bilder, 9)

    # --- Anhang A
    def fill_after(header_start: str, text: str):
        h = _para_after(doc, header_start)
        set_text(_following(doc, h), text)

    fill_after("A3 Lebensräume", ", ".join(lists.LEBENSRAEUME))
    fill_after("A4 Tags", ", ".join(lists.TAGS) + ". (parkplatz-nahe und wc-am-weg setzt die App selbst.)")
    fill_after("A5 Erlebnisse", ", ".join(lists.ERLEBNISSE))
    fill_after("A6 Lizenzen", ", ".join(lists.BILD_LIZENZEN)
               + ". eigen = eigenes Foto, freigabe = schriftliche Freigabe (Bild bleibt unsichtbar, bis sie im CMS bestätigt ist).")

    a7 = _para_after(doc, "A7 Sonstige")
    # A8 nach dem Text von A7 einfuegen
    a7_text = _following(doc, a7)
    for text, bold_src in (("A8 Geometrie-Lizenzen (import.geometrie_lizenz)", a7),
                           (", ".join(lists.GEOMETRIE_LIZENZEN)
                            + ". Ohne gültige Lizenz verwirft der Import die Datei und nimmt die OSM-Pipeline.", a7_text)):
        el = copy.deepcopy(bold_src._p)
        a7_text._p.addnext(el)
        newp = Paragraph(el, doc)
        set_text(newp, text)
        a7_text = newp

    # --- Anhang B
    b_head = _para_after(doc, "Anhang B")
    last = b_head
    el = b_head._p
    # letzten Aufzaehlungspunkt finden
    for p in doc.paragraphs:
        if p.text.startswith("Feste Wertelisten in Anhang A"):
            last = p
    for text in (
        "Version 3: Land, Bundesland, Kreis, Ort, Schutzgebiet haben eigene Spalten",
        "Version 3: Geometrie-Quelle und -Lizenz (Abschnitt 3, Anhang A8)",
        "Version 3: Steckbrief als Tabelle, Geo-Check-Spalte bei Bildern entfällt, Anhang A3-A6 vollständig",
    ):
        el = copy.deepcopy(last._p)
        last._p.addnext(el)
        newp = Paragraph(el, doc)
        set_text(newp, text)
        last = newp

    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)


if __name__ == "__main__":

    if len(sys.argv) < 2:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT)
    print("OK:", sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT)
